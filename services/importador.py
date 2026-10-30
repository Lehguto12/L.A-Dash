"""Leitura, validação e gravação das planilhas enviadas pelos usuários.

Fluxo:  ler_planilha()  ->  validar_dataframe()  ->  consolidar_duplicadas()
        ->  gravar_registros()  (tudo dentro de importar_arquivo()).

Este módulo NÃO conhece Flask nem a sessão: quem chama (a rota) é que decide
qual é o usuario_id e o passa como argumento. Assim ele serve também para a
Fase 2 (Google Sheets), que usará importar_dataframe().
"""
from __future__ import annotations

import io
import os
import unicodedata
import zipfile
from dataclasses import dataclass

import numpy as np
import pandas as pd
from sqlalchemy import delete, insert, select, update
from sqlalchemy.exc import IntegrityError

from extensions import db
from models import Registro

COLUNAS_TEXTO = ("cliente", "campanha", "plataforma")
COLUNAS_INTEIRAS = ("impressoes", "cliques", "leads", "vendas")
COLUNAS_DECIMAIS = ("investimento", "faturamento")
COLUNAS_OBRIGATORIAS = (
    "data", "cliente", "campanha", "plataforma",
    "investimento", "impressoes", "cliques", "leads", "vendas", "faturamento",
)
COLUNAS_METRICAS = COLUNAS_DECIMAIS + COLUNAS_INTEIRAS

EXTENSOES_PERMITIDAS = {".csv", ".xlsx"}
TAMANHO_MAX_TEXTO = 150            # igual ao String(150) das colunas no banco
LIMITE_INTEIRO = 2_000_000_000     # cabe no INTEGER do PostgreSQL (2,147,483,647)
LIMITE_DECIMAL = 1_000_000_000_000
MAX_EXIBIR_ERROS = 10
MAX_BYTES_XLSX_DESCOMPRIMIDO = 100 * 1024 * 1024   # defesa contra "zip bomb"


class ErroImportacao(Exception):
    """Erro que PODE ser mostrado ao usuário (mensagens em português)."""

    def __init__(self, mensagens):
        if isinstance(mensagens, str):
            mensagens = [mensagens]
        self.mensagens = list(mensagens)
        super().__init__("; ".join(self.mensagens))


@dataclass
class ResultadoImportacao:
    inseridos: int = 0
    atualizados: int = 0
    removidos: int = 0
    ignoradas: int = 0
    mescladas: int = 0

    def resumo(self) -> str:
        partes = [f"{self.inseridos} linhas novas"]
        if self.atualizados:
            partes.append(f"{self.atualizados} atualizadas")
        if self.removidos:
            partes.append(f"{self.removidos} antigas removidas")
        if self.mescladas:
            partes.append(f"{self.mescladas} repetidas somadas")
        if self.ignoradas:
            partes.append(f"{self.ignoradas} inválidas ignoradas")
        return "Importação concluída: " + ", ".join(partes) + "."


# ---------------------------------------------------------------- leitura ---

def _normalizar_nome(nome) -> str:
    """'  Impressões ' -> 'impressoes' (sem acento, minúsculo, sem BOM)."""
    texto = str(nome).replace("\ufeff", "").strip().lower()
    texto = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in texto if not unicodedata.combining(c))


def _ler_csv(conteudo: bytes, max_linhas: int) -> pd.DataFrame:
    if b"\x00" in conteudo[:4096]:
        raise ErroImportacao("O arquivo não parece ser um CSV de texto.")
    try:
        texto = conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = conteudo.decode("cp1252", errors="replace")  # Excel antigo no Windows

    # O Excel brasileiro salva CSV com ";" e o padrão internacional usa ",".
    # Olhamos a 1ª linha (cabeçalho) e escolhemos o que aparece mais.
    primeira = texto.split("\n", 1)[0]
    sep = max((";", ",", "\t"), key=primeira.count)
    try:
        return pd.read_csv(
            io.StringIO(texto), sep=sep, dtype=str, keep_default_na=False,
            nrows=max_linhas + 1,
        )
    except (pd.errors.ParserError, pd.errors.EmptyDataError):
        raise ErroImportacao(
            "Não consegui ler o CSV. Confira se todas as linhas têm o mesmo "
            "número de colunas e use o modelo para download."
        )


def _ler_xlsx(conteudo: bytes) -> pd.DataFrame:
    if conteudo[:4] != b"PK\x03\x04":
        raise ErroImportacao("O arquivo não é um .xlsx válido.")
    try:
        with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_BYTES_XLSX_DESCOMPRIMIDO:
                raise ErroImportacao("A planilha é grande demais após descompactada.")
    except zipfile.BadZipFile:
        raise ErroImportacao("O arquivo .xlsx está corrompido.")
    try:
        return pd.read_excel(io.BytesIO(conteudo), dtype=str, engine="openpyxl", sheet_name=0)
    except Exception:
        raise ErroImportacao("Não consegui ler a planilha Excel. Use o modelo para download.")


def ler_planilha(arquivo, max_linhas: int) -> pd.DataFrame:
    """Lê um arquivo enviado (FileStorage) direto da memória.

    Nunca gravamos o arquivo em disco, então o nome enviado pelo usuário só é
    usado para descobrir a extensão (nada de path traversal)."""
    extensao = os.path.splitext(arquivo.filename or "")[1].lower()
    if extensao not in EXTENSOES_PERMITIDAS:
        raise ErroImportacao("Formato não suportado. Envie um arquivo .csv ou .xlsx.")

    conteudo = arquivo.read()
    if not conteudo:
        raise ErroImportacao("O arquivo está vazio.")

    df = _ler_csv(conteudo, max_linhas) if extensao == ".csv" else _ler_xlsx(conteudo)
    if len(df) > max_linhas:
        raise ErroImportacao(f"O arquivo tem mais de {max_linhas:,} linhas (limite por envio).".replace(",", "."))
    return df


# ------------------------------------------------------------- validação ---

def _preparar_colunas(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = [_normalizar_nome(c) for c in df.columns]
    if df.columns.duplicated().any():
        raise ErroImportacao("Há colunas com o mesmo nome na planilha.")
    faltando = [c for c in COLUNAS_OBRIGATORIAS if c not in df.columns]
    if faltando:
        raise ErroImportacao([
            "Faltam colunas obrigatórias: " + ", ".join(faltando) + ".",
            "Colunas esperadas: " + ", ".join(COLUNAS_OBRIGATORIAS) + ".",
        ])
    df = df[list(COLUNAS_OBRIGATORIAS)].fillna("")
    return df.apply(lambda col: col.astype(str).str.strip())


def _para_numero(serie: pd.Series) -> pd.Series:
    """Aceita '1234.5', '1.234,50', '1.000' e 'R$ 1.234,50'. Vazio vira 0.

    Regras: se há vírgula, é formato brasileiro (ponto = milhar, vírgula =
    decimal). Sem vírgula, um número como 1.000 ou 12.345.678 (grupos de 3
    dígitos após o ponto, sem começar por 0) é tratado como milhar; qualquer
    outro (350.5, 0.500) é decimal com ponto."""
    s = serie.str.replace(r"[Rr]\$", "", regex=True).str.replace(r"\s", "", regex=True)
    br = s.str.contains(",", regex=False)          # tem vírgula => formato brasileiro
    s = s.where(~br, s.str.replace(".", "", regex=False).str.replace(",", ".", regex=False))
    milhar = s.str.fullmatch(r"[1-9]\d{0,2}(\.\d{3})+")
    s = s.where(~milhar, s.str.replace(".", "", regex=False))
    s = s.replace("", "0")
    return pd.to_numeric(s, errors="coerce")


def _para_data(serie: pd.Series) -> pd.Series:
    """Converte texto em data SEM ambiguidade.

    Tentamos primeiro o padrão ISO (AAAA-MM-DD, com ou sem hora, como o Excel
    entrega) e só depois os formatos brasileiros (DD/MM/AAAA). Não usamos
    dayfirst=True direto porque ele embaralha datas ISO (2026-10-01 viraria
    10 de janeiro)."""
    datas = pd.to_datetime(serie, format="ISO8601", errors="coerce")
    for formato in ("%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y"):
        faltam = datas.isna()
        if not faltam.any():
            break
        datas[faltam] = pd.to_datetime(serie[faltam], format=formato, errors="coerce")
    return datas


def validar_dataframe(df: pd.DataFrame):
    """Devolve (df_valido, mensagens_de_erro, qtd_linhas_invalidas)."""
    df = _preparar_colunas(df)
    motivos: dict[int, list[str]] = {}

    def marcar(mascara: pd.Series, texto: str):
        for idx in mascara[mascara].index:
            motivos.setdefault(idx, []).append(texto)

    for col in COLUNAS_TEXTO:
        marcar(df[col] == "", f"'{col}' vazio")
        marcar(df[col].str.len() > TAMANHO_MAX_TEXTO, f"'{col}' com mais de {TAMANHO_MAX_TEXTO} caracteres")

    datas = _para_data(df["data"])
    marcar(datas.isna() | (datas < "2000-01-01") | (datas > "2100-12-31"),
           "'data' inválida (use AAAA-MM-DD ou DD/MM/AAAA)")

    numeros = {}
    for col in COLUNAS_METRICAS:
        v = _para_numero(df[col])
        numeros[col] = v
        marcar(v.isna() | np.isinf(v), f"'{col}' não é um número")
        marcar(v < 0, f"'{col}' não pode ser negativo")
        limite = LIMITE_INTEIRO if col in COLUNAS_INTEIRAS else LIMITE_DECIMAL
        marcar(v > limite, f"'{col}' grande demais")
        if col in COLUNAS_INTEIRAS:
            marcar(v.notna() & np.isfinite(v) & (v % 1 != 0), f"'{col}' deve ser inteiro")

    invalidas = pd.Index(sorted(motivos))
    # Linha no arquivo = índice + 2 (a linha 1 é o cabeçalho).
    erros = [f"Linha {i + 2}: " + "; ".join(motivos[i]) for i in invalidas[:MAX_EXIBIR_ERROS]]
    if len(invalidas) > MAX_EXIBIR_ERROS:
        erros.append(f"... e mais {len(invalidas) - MAX_EXIBIR_ERROS} linhas com problema.")

    ok = ~df.index.isin(invalidas)
    limpo = pd.DataFrame({
        "data": datas[ok].dt.date,
        **{c: df.loc[ok, c] for c in COLUNAS_TEXTO},
    })
    for col in COLUNAS_DECIMAIS:
        limpo[col] = numeros[col][ok].astype(float)
    for col in COLUNAS_INTEIRAS:
        limpo[col] = numeros[col][ok].astype("int64")
    return limpo[list(COLUNAS_OBRIGATORIAS)].reset_index(drop=True), erros, len(invalidas)


def consolidar_duplicadas(df: pd.DataFrame):
    """Mesma chave (data+cliente+plataforma+campanha) repetida: soma as métricas,
    porque o banco só aceita UMA linha por chave (restrição única)."""
    chave = ["data", "cliente", "plataforma", "campanha"]
    agrupado = df.groupby(chave, as_index=False, sort=False)[list(COLUNAS_METRICAS)].sum()
    return agrupado[list(COLUNAS_OBRIGATORIAS)], len(df) - len(agrupado)


# --------------------------------------------------------------- gravação ---

def gravar_registros(df: pd.DataFrame, usuario_id: int, modo: str) -> ResultadoImportacao:
    """Grava numa ÚNICA transação: ou entra tudo, ou nada muda no banco."""
    linhas = df.to_dict("records")
    for linha in linhas:
        linha["usuario_id"] = usuario_id

    resultado = ResultadoImportacao()
    try:
        if modo == "substituir":
            r = db.session.execute(delete(Registro).where(Registro.usuario_id == usuario_id))
            resultado.removidos = r.rowcount or 0
            novos, existentes_atualizar = linhas, []
        else:  # "atualizar"
            # Só buscamos ids de linhas DESTE usuário, então o UPDATE abaixo
            # jamais encosta em dados de outra pessoa.
            ids = {
                (d, c, p, ca): i
                for i, d, c, p, ca in db.session.execute(
                    select(Registro.id, Registro.data, Registro.cliente,
                           Registro.plataforma, Registro.campanha)
                    .where(Registro.usuario_id == usuario_id)
                )
            }
            novos, existentes_atualizar = [], []
            for linha in linhas:
                chave = (linha["data"], linha["cliente"], linha["plataforma"], linha["campanha"])
                if chave in ids:
                    existentes_atualizar.append(
                        {"id": ids[chave], **{m: linha[m] for m in COLUNAS_METRICAS}}
                    )
                else:
                    novos.append(linha)

        if existentes_atualizar:
            db.session.execute(update(Registro), existentes_atualizar)
        if novos:
            db.session.execute(insert(Registro), novos)
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        raise ErroImportacao("Conflito ao gravar (dois envios ao mesmo tempo?). Tente novamente.")
    except Exception:
        db.session.rollback()
        raise

    resultado.inseridos = len(novos)
    resultado.atualizados = len(existentes_atualizar)
    return resultado


def importar_dataframe(df, usuario_id: int, modo: str = "atualizar",
                       ignorar_invalidas: bool = False) -> ResultadoImportacao:
    limpo, erros, qtd_invalidas = validar_dataframe(df)
    if qtd_invalidas and not ignorar_invalidas:
        raise ErroImportacao(
            ["Nenhum dado foi importado. Corrija as linhas abaixo e envie de novo:"] + erros
        )
    if limpo.empty:
        raise ErroImportacao(["Nenhuma linha válida encontrada."] + erros)

    limpo, mescladas = consolidar_duplicadas(limpo)
    resultado = gravar_registros(limpo, usuario_id, modo)
    resultado.ignoradas = qtd_invalidas
    resultado.mescladas = mescladas
    return resultado


def importar_arquivo(arquivo, usuario_id: int, modo: str = "atualizar",
                     ignorar_invalidas: bool = False, max_linhas: int = 50_000) -> ResultadoImportacao:
    df = ler_planilha(arquivo, max_linhas)
    return importar_dataframe(df, usuario_id, modo, ignorar_invalidas)


# ----------------------------------------------------------------- modelo ---

_EXEMPLOS = [
    ("2026-10-01", "Cliente Exemplo", "Conversao_Remarketing", "Meta Ads", 350.50, 48200, 1250, 96, 8, 4320.00),
    ("2026-10-01", "Cliente Exemplo", "Search_Institucional", "Google Ads", 210.00, 15300, 640, 41, 5, 2700.00),
    ("2026-10-02", "Cliente Exemplo", "Conversao_Remarketing", "Meta Ads", 340.00, 46900, 1190, 88, 7, 3780.00),
]
_DESCRICOES = {
    "data": "Dia da campanha. Aceita AAAA-MM-DD ou DD/MM/AAAA.",
    "cliente": "Nome do cliente. Cada nome diferente conta como 1 cliente no seu plano.",
    "campanha": "Nome da campanha.",
    "plataforma": "Ex.: Meta Ads, Google Ads, TikTok Ads.",
    "investimento": "Valor investido no dia (R$). Aceita 350.50 ou 350,50.",
    "impressoes": "Número inteiro.",
    "cliques": "Número inteiro.",
    "leads": "Número inteiro.",
    "vendas": "Número inteiro.",
    "faturamento": "Receita gerada no dia (R$).",
}


def gerar_modelo_csv() -> bytes:
    """CSV no padrão do Excel brasileiro: ';' como separador e ',' decimal."""
    linhas = [";".join(COLUNAS_OBRIGATORIAS)]
    for ex in _EXEMPLOS:
        linhas.append(";".join(
            f"{v:.2f}".replace(".", ",") if isinstance(v, float) else str(v) for v in ex
        ))
    return ("\r\n".join(linhas) + "\r\n").encode("utf-8-sig")


def gerar_modelo_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "dados"
    ws.append(list(COLUNAS_OBRIGATORIAS))
    for ex in _EXEMPLOS:
        ws.append(list(ex))
    for cel in ws[1]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="7C3AED")
    for letra in "ABCDEFGHIJ":
        ws.column_dimensions[letra].width = 20

    guia = wb.create_sheet("instrucoes")
    guia.append(["coluna", "como preencher"])
    for col in COLUNAS_OBRIGATORIAS:
        guia.append([col, _DESCRICOES[col]])
    guia.column_dimensions["A"].width = 16
    guia.column_dimensions["B"].width = 80
    guia["A1"].font = guia["B1"].font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()